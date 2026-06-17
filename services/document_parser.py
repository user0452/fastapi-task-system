from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_DOCUMENT_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}


def _decode_text_file(content: bytes) -> str:
    """
    解析 txt / md 文本文件。
    优先 utf-8，失败后尝试 gbk。
    """
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("gbk", errors="ignore")


def _extract_pdf_text(content: bytes) -> str:
    """
    解析文本型 PDF。
    注意：扫描版 PDF / 图片型 PDF 不支持 OCR。
    """
    reader = PdfReader(BytesIO(content))

    texts = []

    for index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""

        if page_text.strip():
            texts.append(f"【第 {index} 页】\n{page_text.strip()}")

    return "\n\n".join(texts)


def _extract_docx_text(content: bytes) -> str:
    """
    解析 docx 文档。
    支持段落和简单表格文本。
    """
    document = Document(BytesIO(content))

    lines = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    cells.append(cell_text)

            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)


def extract_text_from_document(filename: str, content: bytes) -> str:
    """
    根据文件后缀解析课程资料文本。
    """
    suffix = Path(filename).suffix.lower()

    if suffix not in SUPPORTED_DOCUMENT_SUFFIXES:
        raise ValueError(
            f"暂不支持该文件类型：{suffix}，仅支持 txt、md、pdf、docx"
        )

    if suffix in {".txt", ".md"}:
        text = _decode_text_file(content)

    elif suffix == ".pdf":
        text = _extract_pdf_text(content)

    elif suffix == ".docx":
        text = _extract_docx_text(content)

    else:
        raise ValueError(f"暂不支持该文件类型：{suffix}")

    text = text.replace("\ufeff", "").strip()

    if not text:
        raise ValueError(
            "未能从文件中解析出文本内容。如果是扫描版 PDF 或图片型 PDF，当前版本暂不支持 OCR。"
        )

    return text