from io import BytesIO

from docx import Document

from app.integrations.document_parser import extract_text_from_document
from app.integrations.embedding.chunking import chunk_document, retrieval_text


def test_structured_chunking_preserves_heading_kb_id_and_token_budget():
    text = """# Product manual

## Access control

### Export policy

**KB-ID：YQ-SEC-052**

Raw audit exports require a separate compliance authorization.

### Backup policy

**KB-ID: YQ-DR-002**

Backup success does not prove restore success. Restore verification is mandatory.
"""

    chunks = chunk_document(text)

    assert {kb_id for chunk in chunks for kb_id in chunk["kb_ids"]} == {
        "YQ-SEC-052",
        "YQ-DR-002",
    }
    assert all(len(chunk["kb_ids"]) == 1 for chunk in chunks if chunk["kb_ids"])
    assert all(chunk["estimated_tokens"] <= 320 for chunk in chunks)
    export = next(chunk for chunk in chunks if "YQ-SEC-052" in chunk["kb_ids"])
    assert export["heading_path"].endswith("Access control > Export policy")
    assert text[export["char_start"] : export["char_end"]] == export["chunk_text"]
    assert "YQ-SEC-052" in retrieval_text(
        export["chunk_text"], export["heading_path"], export["kb_ids"], "Manual"
    )


def test_page_marker_is_kept_as_chunk_metadata():
    chunks = chunk_document("【第 2 页】\n\n**KB-ID：YQ-PAGE-002**\n\nPage scoped fact.")

    assert len(chunks) == 1
    assert chunks[0]["page_number"] == 2
    assert chunks[0]["kb_ids"] == ["YQ-PAGE-002"]


def test_docx_parser_preserves_paragraph_table_paragraph_order():
    document = Document()
    document.add_heading("Runbook", level=1)
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "field"
    table.cell(0, 1).text = "meaning"
    document.add_paragraph("After table")
    buffer = BytesIO()
    document.save(buffer)

    parsed = extract_text_from_document("runbook.docx", buffer.getvalue())

    assert parsed.index("# Runbook") < parsed.index("Before table")
    assert parsed.index("Before table") < parsed.index("| field | meaning |")
    assert parsed.index("| field | meaning |") < parsed.index("After table")
