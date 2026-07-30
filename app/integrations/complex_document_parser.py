"""Layout-aware TXT, Markdown, PDF and DOCX parsing with optional OCR."""

from __future__ import annotations

import math
import os
import re
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any, Callable
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.integrations.document_model import DocumentBlock, ParsedDocument

SUPPORTED_DOCUMENT_SUFFIXES = {".txt", ".md", ".pdf", ".docx"}
MAX_DOCUMENT_PAGES = 2_000
MAX_EXTRACTED_TEXT_CHARS = 12_000_000
MAX_DOCX_UNCOMPRESSED_SIZE = 256 * 1024 * 1024
MAX_DOCX_MEMBERS = 20_000
OCR_MIN_PAGE_TEXT_CHARS = int(os.getenv("RAG_OCR_MIN_PAGE_TEXT_CHARS", "24"))
OCR_LANGUAGE = os.getenv("RAG_OCR_LANGUAGE", "chi_sim+eng").strip() or "eng"
OCR_DPI = max(96, min(int(os.getenv("RAG_OCR_DPI", "200")), 400))
OCR_TESSDATA = os.getenv("RAG_OCR_TESSDATA", "").strip() or None

_PAGE_NUMBER = re.compile(r"^(?:page\s*)?\d+(?:\s*[/\-]\s*\d+)?$", re.IGNORECASE)
_MARKDOWN_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")


def _ensure_text_size(length: int) -> None:
    if length > MAX_EXTRACTED_TEXT_CHARS:
        raise ValueError(
            f"资料解析文本超过 {MAX_EXTRACTED_TEXT_CHARS // 1_000_000}M 字符，"
            "请拆分后上传，避免索引任务占用过多内存"
        )


def _validate_suffix(filename: str | Path) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_SUFFIXES:
        raise ValueError(f"暂不支持 {suffix or '无扩展名'} 文件，仅支持 txt、md、pdf、docx")
    return suffix


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("gbk", errors="ignore")


def _read_text_path(source: Path) -> str:
    with source.open("rb") as stream:
        prefix = stream.read(4)
    if b"\x00" in prefix:
        raise ValueError("文本文件包含二进制内容，已拒绝解析")
    try:
        with source.open("r", encoding="utf-8-sig") as stream:
            return stream.read(MAX_EXTRACTED_TEXT_CHARS + 1)
    except UnicodeDecodeError:
        with source.open("r", encoding="gbk", errors="ignore") as stream:
            return stream.read(MAX_EXTRACTED_TEXT_CHARS + 1)


def _text_blocks(source_name: str, text: str) -> ParsedDocument:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").replace("\ufeff", "").strip()
    _ensure_text_size(len(normalized))
    if not normalized:
        raise ValueError("资料内容为空，无法构建索引")
    blocks: list[DocumentBlock] = []
    for raw in re.split(r"\n\s*\n", normalized):
        value = raw.strip()
        if not value:
            continue
        heading = _MARKDOWN_HEADING.match(value)
        block_type = "heading" if heading else "paragraph"
        content = heading.group(2).strip() if heading else value
        blocks.append(
            DocumentBlock(
                block_id=f"b{len(blocks) + 1}",
                block_index=len(blocks),
                reading_order=len(blocks),
                block_type=block_type,
                text=content,
                heading_level=len(heading.group(1)) if heading else None,
            )
        )
    return ParsedDocument(source_name=source_name, blocks=blocks)


def _docx_paragraph_block(paragraph: Paragraph, index: int) -> DocumentBlock | None:
    text = paragraph.text.strip()
    if not text:
        return None
    style_name = str(paragraph.style.name if paragraph.style else "").strip()
    heading = re.match(r"Heading\s+(\d+)$", style_name, re.IGNORECASE)
    if heading or style_name.lower() == "title":
        return DocumentBlock(
            block_id=f"b{index + 1}",
            block_index=index,
            reading_order=index,
            block_type="heading",
            text=text,
            heading_level=max(1, min(6, int(heading.group(1)))) if heading else 1,
            metadata={"docx_style": style_name},
        )
    block_type = "list" if "list" in style_name.lower() else "paragraph"
    return DocumentBlock(
        block_id=f"b{index + 1}",
        block_index=index,
        reading_order=index,
        block_type=block_type,
        text=text,
        metadata={"docx_style": style_name} if style_name else {},
    )


def _docx_table_block(table: Table, index: int) -> DocumentBlock | None:
    cells = [
        [cell.text.strip().replace("\n", " ") for cell in row.cells]
        for row in table.rows
    ]
    if not any(any(cell for cell in row) for row in cells):
        return None
    return DocumentBlock(
        block_id=f"b{index + 1}",
        block_index=index,
        reading_order=index,
        block_type="table",
        table_cells=cells,
        metadata={"rows": len(cells), "columns": max((len(row) for row in cells), default=0)},
    )


def _docx_noise_blocks(document: DocumentType, start_index: int) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    seen: set[tuple[str, str]] = set()
    for section in document.sections:
        for label, container in (("repeated_header", section.header), ("repeated_footer", section.footer)):
            text = "\n".join(paragraph.text.strip() for paragraph in container.paragraphs if paragraph.text.strip())
            key = (label, text)
            if not text or key in seen:
                continue
            seen.add(key)
            index = start_index + len(blocks)
            blocks.append(
                DocumentBlock(
                    block_id=f"b{index + 1}",
                    block_index=index,
                    reading_order=index,
                    block_type="noise",
                    text=text,
                    should_index=False,
                    noise_reason=label,
                )
            )
    return blocks


def _parse_docx_document(source_name: str, document: DocumentType) -> ParsedDocument:
    blocks: list[DocumentBlock] = []
    indexed_chars = 0
    for item in document.iter_inner_content():
        block = (
            _docx_paragraph_block(item, len(blocks))
            if isinstance(item, Paragraph)
            else _docx_table_block(item, len(blocks))
            if isinstance(item, Table)
            else None
        )
        if block is not None:
            blocks.append(block)
            indexed_chars += len(block.index_text()) + 2
            _ensure_text_size(indexed_chars)
    blocks.extend(_docx_noise_blocks(document, len(blocks)))
    if not any(block.should_index and block.index_text() for block in blocks):
        raise ValueError("未能从 DOCX 中解析出可索引内容")
    parsed = ParsedDocument(source_name=source_name, blocks=blocks)
    _ensure_text_size(len(parsed.render_for_index().text))
    return parsed


def _rect_overlap_ratio(left: tuple[float, float, float, float], right: tuple[float, float, float, float]) -> float:
    x0, y0 = max(left[0], right[0]), max(left[1], right[1])
    x1, y1 = min(left[2], right[2]), min(left[3], right[3])
    intersection = max(0.0, x1 - x0) * max(0.0, y1 - y0)
    area = max(1.0, (left[2] - left[0]) * (left[3] - left[1]))
    return intersection / area


def _dict_text_block(raw: dict[str, Any]) -> str:
    lines = []
    for line in raw.get("lines") or []:
        content = "".join(str(span.get("text") or "") for span in line.get("spans") or []).strip()
        if content:
            lines.append(content)
    return "\n".join(lines).strip()


def _bbox(values: Any) -> tuple[float, float, float, float]:
    x0, y0, x1, y1 = values
    return float(x0), float(y0), float(x1), float(y1)


def _ocr_page_dict(page) -> dict[str, Any]:
    textpage = page.get_textpage_ocr(
        language=OCR_LANGUAGE,
        dpi=OCR_DPI,
        full=True,
        tessdata=OCR_TESSDATA,
    )
    return page.get_text("dict", textpage=textpage, sort=True)


def _extract_pdf_page(
    page,
    page_number: int,
    start_index: int,
    ocr_provider: Callable[[Any], dict[str, Any]] | None,
) -> tuple[list[DocumentBlock], list[str]]:
    warnings: list[str] = []
    native_page_dict = page.get_text("dict", sort=True)
    page_dict = native_page_dict
    plain_text = "".join(
        _dict_text_block(block)
        for block in page_dict.get("blocks") or []
        if int(block.get("type", 0)) == 0
    )
    ocr_used = False
    if len(re.sub(r"\s+", "", plain_text)) < OCR_MIN_PAGE_TEXT_CHARS:
        try:
            ocr_page_dict = (ocr_provider or _ocr_page_dict)(page)
            native_images = [
                block
                for block in native_page_dict.get("blocks") or []
                if int(block.get("type", 0)) == 1
            ]
            page_dict = {
                **ocr_page_dict,
                "blocks": [*(ocr_page_dict.get("blocks") or []), *native_images],
            }
            ocr_used = True
        except Exception as exc:
            warnings.append(f"第 {page_number} 页需要 OCR，但 OCR 不可用：{type(exc).__name__}")

    table_payloads: list[tuple[tuple[float, float, float, float], list[list[str]]]] = []
    try:
        finder = page.find_tables()
        for table in finder.tables:
            cells = [["" if cell is None else str(cell) for cell in row] for row in table.extract()]
            if any(any(cell.strip() for cell in row) for row in cells):
                table_payloads.append((_bbox(table.bbox), cells))
    except Exception as exc:
        warnings.append(f"第 {page_number} 页表格检测失败：{type(exc).__name__}")

    pending: list[tuple[float, float, str, str, tuple[float, float, float, float], Any]] = []
    for bbox, cells in table_payloads:
        pending.append((bbox[1], bbox[0], "table", "", bbox, cells))
    for raw in page_dict.get("blocks") or []:
        bbox = _bbox(raw.get("bbox", (0, 0, 0, 0)))
        raw_type = int(raw.get("type", 0))
        if raw_type == 1:
            pending.append((bbox[1], bbox[0], "image", "", bbox, raw))
            continue
        if raw_type != 0 or any(_rect_overlap_ratio(bbox, table_bbox) >= 0.5 for table_bbox, _ in table_payloads):
            continue
        text = _dict_text_block(raw)
        if text:
            pending.append((bbox[1], bbox[0], "paragraph", text, bbox, raw))

    pending.sort(key=lambda item: (round(item[0], 1), round(item[1], 1), item[2]))
    blocks: list[DocumentBlock] = []
    for order, (_y, _x, block_type, text, bbox, payload) in enumerate(pending):
        index = start_index + len(blocks)
        if block_type == "table":
            block = DocumentBlock(
                block_id=f"p{page_number}-b{order + 1}",
                block_index=index,
                reading_order=order,
                page_number=page_number,
                bbox=bbox,
                block_type="table",
                table_cells=payload,
                ocr_used=ocr_used,
            )
        elif block_type == "image":
            block = DocumentBlock(
                block_id=f"p{page_number}-b{order + 1}",
                block_index=index,
                reading_order=order,
                page_number=page_number,
                bbox=bbox,
                block_type="image",
                should_index=False,
                noise_reason="image_not_transcribed",
                metadata={
                    "width": payload.get("width"),
                    "height": payload.get("height"),
                },
            )
        else:
            block = DocumentBlock(
                block_id=f"p{page_number}-b{order + 1}",
                block_index=index,
                reading_order=order,
                page_number=page_number,
                bbox=bbox,
                block_type="paragraph",
                text=text,
                ocr_used=ocr_used,
            )
        blocks.append(block)
    return blocks, warnings


def _normalize_margin_text(text: str) -> str:
    return re.sub(r"\s+", "", text).casefold()


def _mark_repeated_margins(blocks: list[DocumentBlock], page_heights: dict[int, float]) -> None:
    pages = set(page_heights)
    threshold = max(2, math.ceil(len(pages) * 0.6))
    occurrences: dict[tuple[str, str], set[int]] = defaultdict(set)
    for block in blocks:
        if not block.text or block.page_number is None or block.bbox is None:
            continue
        height = page_heights[block.page_number]
        region = "header" if block.bbox[1] <= height * 0.12 else "footer" if block.bbox[3] >= height * 0.88 else ""
        if not region:
            continue
        normalized = _normalize_margin_text(block.text)
        if normalized:
            occurrences[(region, normalized)].add(block.page_number)

    repeated = {key for key, page_numbers in occurrences.items() if len(page_numbers) >= threshold}
    for block in blocks:
        if not block.text or block.page_number is None or block.bbox is None:
            continue
        height = page_heights[block.page_number]
        region = "header" if block.bbox[1] <= height * 0.12 else "footer" if block.bbox[3] >= height * 0.88 else ""
        normalized = _normalize_margin_text(block.text)
        reason = None
        if region and (region, normalized) in repeated:
            reason = f"repeated_{region}"
        elif region and _PAGE_NUMBER.fullmatch(normalized):
            reason = "page_number"
        if reason:
            block.should_index = False
            block.noise_reason = reason
            block.block_type = "noise"


def _parse_pdf_document(
    source_name: str,
    document,
    *,
    ocr_provider: Callable[[Any], dict[str, Any]] | None = None,
) -> ParsedDocument:
    if len(document) > MAX_DOCUMENT_PAGES:
        raise ValueError(f"PDF 页数超过 {MAX_DOCUMENT_PAGES} 页，请拆分后上传")
    blocks: list[DocumentBlock] = []
    warnings: list[str] = []
    page_heights: dict[int, float] = {}
    indexed_chars = 0
    for offset, page in enumerate(document):
        page_number = offset + 1
        page_heights[page_number] = float(page.rect.height)
        page_blocks, page_warnings = _extract_pdf_page(
            page,
            page_number,
            len(blocks),
            ocr_provider,
        )
        blocks.extend(page_blocks)
        warnings.extend(page_warnings)
        indexed_chars += sum(len(block.index_text()) + 2 for block in page_blocks)
        _ensure_text_size(indexed_chars)
    _mark_repeated_margins(blocks, page_heights)
    parsed = ParsedDocument(source_name=source_name, blocks=blocks, warnings=warnings)
    rendered = parsed.render_for_index().text
    _ensure_text_size(len(rendered))
    if not rendered:
        detail = "；".join(warnings[:3])
        raise ValueError(
            "未能从 PDF 中解析出可索引文本；扫描版需要安装 Tesseract 和中文语言包"
            + (f"（{detail}）" if detail else "")
        )
    return parsed


def _validate_docx_members(archive: ZipFile) -> None:
    members = archive.infolist()
    if len(members) > MAX_DOCX_MEMBERS:
        raise ValueError("DOCX 内部文件数量异常，已拒绝解析")
    if sum(item.file_size for item in members) > MAX_DOCX_UNCOMPRESSED_SIZE:
        raise ValueError("DOCX 解压后体积过大，请拆分后上传")
    names = {item.filename for item in members}
    if "[Content_Types].xml" not in names or "word/document.xml" not in names:
        raise ValueError("文件内容不是有效的 DOCX 文档")


def _validate_docx_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            _validate_docx_members(archive)
    except BadZipFile as exc:
        raise ValueError("文件内容不是有效的 DOCX 文档") from exc


def _validate_docx_bytes(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            _validate_docx_members(archive)
    except BadZipFile as exc:
        raise ValueError("文件内容不是有效的 DOCX 文档") from exc


def parse_document_from_bytes(
    filename: str,
    content: bytes,
    *,
    ocr_provider: Callable[[Any], dict[str, Any]] | None = None,
) -> ParsedDocument:
    suffix = _validate_suffix(filename)
    if suffix in {".txt", ".md"}:
        return _text_blocks(filename, _decode_text(content))
    if suffix == ".docx":
        _validate_docx_bytes(content)
        return _parse_docx_document(filename, Document(BytesIO(content)))
    try:
        with pymupdf.open(stream=content, filetype="pdf") as document:
            return _parse_pdf_document(filename, document, ocr_provider=ocr_provider)
    except pymupdf.FileDataError as exc:
        raise ValueError("文件内容不是有效的 PDF 文档") from exc


def parse_document_from_path(
    path: str | Path,
    filename: str | None = None,
    *,
    ocr_provider: Callable[[Any], dict[str, Any]] | None = None,
) -> ParsedDocument:
    source = Path(path)
    source_name = filename or source.name
    suffix = _validate_suffix(source_name)
    if suffix in {".txt", ".md"}:
        return _text_blocks(source_name, _read_text_path(source))
    if suffix == ".docx":
        _validate_docx_archive(source)
        return _parse_docx_document(source_name, Document(str(source)))
    with source.open("rb") as stream:
        if stream.read(5) != b"%PDF-":
            raise ValueError("文件内容不是有效的 PDF 文档")
    try:
        with pymupdf.open(str(source)) as document:
            return _parse_pdf_document(source_name, document, ocr_provider=ocr_provider)
    except pymupdf.FileDataError as exc:
        raise ValueError("文件内容不是有效的 PDF 文档") from exc


def parse_text_document(source_name: str, text: str) -> ParsedDocument:
    return _text_blocks(source_name, text)


__all__ = [
    "MAX_DOCUMENT_PAGES",
    "MAX_DOCX_MEMBERS",
    "MAX_DOCX_UNCOMPRESSED_SIZE",
    "MAX_EXTRACTED_TEXT_CHARS",
    "OCR_DPI",
    "OCR_LANGUAGE",
    "SUPPORTED_DOCUMENT_SUFFIXES",
    "parse_document_from_bytes",
    "parse_document_from_path",
    "parse_text_document",
]
